#!/usr/bin/env python3
"""Convert RIKEN Lipidomics Excel alignment tables to mzTab-M 2.0-M.

The workbooks contain a variable-length annotation block followed by raw peak
heights, a ``Lipid IS`` separator, and normalized abundances.  This converter
detects those boundaries from labels instead of assuming fixed column numbers.
"""

from __future__ import annotations

import argparse
import math
import re
import sys
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - exercised by CLI environments
    raise SystemExit("openpyxl is required: python -m pip install openpyxl") from exc


NULL = "null"
METADATA_LABELS = {
    "public/private",
    "category",
    "tissue/species",
    "genotype/background",
    "perturbation",
    "diet/culture",
    "biological replicate",
    "technical replicate",
    "unit",
}
HEADER_ANCHORS = {"alignment id", "metabolite name", "average mz"}
SECTION_SEPARATORS = {"lipid is", "standard class"}


@dataclass
class ParsedSheet:
    source: Path
    sheet: str
    header_row: int
    headers: list[str]
    annotation_columns: list[int]
    abundance_columns: list[int]
    sample_names: list[str]
    sample_metadata: list[dict[str, object]]
    rows: list[tuple[object, ...]]
    abundance_kind: str


def clean(value: object, fallback: str = NULL) -> str:
    if value is None:
        return fallback
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return fallback
        value = format(value, ".15g")
    text = re.sub(r"[\t\r\n]+", " ", str(value)).strip()
    return text or fallback


def key(value: object) -> str:
    return re.sub(r"\s+", " ", clean(value, "").lower())


def unique_names(names: Sequence[str]) -> list[str]:
    counts: dict[str, int] = {}
    result: list[str] = []
    for index, original in enumerate(names, 1):
        base = original if original and original != NULL else f"sample_{index}"
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}.{counts[base]}")
    return result


def find_header_row(ws, scan_rows: int = 50) -> int:
    for row_number, row in enumerate(
        ws.iter_rows(min_row=1, max_row=min(ws.max_row, scan_rows), values_only=True), 1
    ):
        present = {key(value) for value in row if value is not None}
        if len(HEADER_ANCHORS & present) >= 2 and "metabolite name" in present:
            return row_number
    raise ValueError(f"{ws.title}: alignment-table header was not found in the first {scan_rows} rows")


def parse_sheet(
    source: Path,
    ws,
    abundance: str = "normalized",
    privacy: str = "public",
) -> ParsedSheet:
    header_row = find_header_row(ws)
    all_rows = list(ws.iter_rows(min_row=1, values_only=True))
    headers = [clean(value, "") for value in all_rows[header_row - 1]]
    lowered = [key(value) for value in headers]

    try:
        spectrum_col = lowered.index("ms/ms spectrum")
    except ValueError as exc:
        raise ValueError(f"{ws.title}: required 'MS/MS spectrum' column is missing") from exc
    separator_col = next(
        (index for index in range(spectrum_col + 1, len(lowered)) if lowered[index] in SECTION_SEPARATORS),
        None,
    )
    if separator_col is None:
        raise ValueError(f"{ws.title}: abundance section separator ('Lipid IS' or 'standard class') is missing")

    raw_columns = list(range(spectrum_col + 1, separator_col))
    normalized_columns = list(range(separator_col + 1, len(headers)))
    selected = normalized_columns if abundance == "normalized" else raw_columns
    if not selected:
        raise ValueError(f"{ws.title}: the {abundance} abundance section is empty")

    label_column = spectrum_col
    metadata_rows: dict[str, int] = {}
    for row_index in range(header_row - 1):
        label = key(all_rows[row_index][label_column] if label_column < len(all_rows[row_index]) else None)
        if label in METADATA_LABELS:
            metadata_rows[label] = row_index

    keep: list[int] = []
    metadata: list[dict[str, object]] = []
    for column in selected:
        record = {
            label: (all_rows[row][column] if column < len(all_rows[row]) else None)
            for label, row in metadata_rows.items()
        }
        visibility = key(record.get("public/private"))
        if privacy == "public" and visibility not in {"", "public"}:
            continue
        keep.append(column)
        metadata.append(record)
    if not keep:
        raise ValueError(f"{ws.title}: no samples remain after --privacy {privacy}")

    sample_names = unique_names([headers[column] for column in keep])
    data_rows = [tuple(row) for row in all_rows[header_row:] if any(value is not None for value in row)]
    return ParsedSheet(
        source=source,
        sheet=ws.title,
        header_row=header_row,
        headers=headers,
        annotation_columns=list(range(0, spectrum_col + 1)),
        abundance_columns=keep,
        sample_names=sample_names,
        sample_metadata=metadata,
        rows=data_rows,
        abundance_kind=abundance,
    )


def parse_workbook(
    source: Path,
    sheet: str | None = None,
    abundance: str = "normalized",
    privacy: str = "public",
) -> list[ParsedSheet]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    worksheets = [workbook[sheet]] if sheet else list(workbook.worksheets)
    parsed: list[ParsedSheet] = []
    errors: list[str] = []
    for ws in worksheets:
        try:
            parsed.append(parse_sheet(source, ws, abundance=abundance, privacy=privacy))
        except ValueError as exc:
            errors.append(str(exc))
    workbook.close()
    if sheet and errors:
        raise ValueError(errors[0])
    if not parsed:
        raise ValueError(f"{source.name}: no convertible sheet found ({'; '.join(errors)})")
    return parsed


def value_at(row: Sequence[object], index: int | None) -> object:
    return None if index is None or index >= len(row) else row[index]


def column_index(parsed: ParsedSheet, *aliases: str) -> int | None:
    normalized = {key(header): index for index, header in enumerate(parsed.headers)}
    for alias in aliases:
        if key(alias) in normalized:
            return normalized[key(alias)]
    return None


def group_label(metadata: dict[str, object], group_by: Sequence[str]) -> str:
    parts = [clean(metadata.get(key(name)), "") for name in group_by]
    parts = [part for part in parts if part]
    return " | ".join(parts) or "Unknown"


def write_mztab(parsed: ParsedSheet, destination: Path, group_by: Sequence[str]) -> None:
    groups: OrderedDict[str, list[int]] = OrderedDict()
    for assay_index, metadata in enumerate(parsed.sample_metadata, 1):
        groups.setdefault(group_label(metadata, group_by), []).append(assay_index)

    lines: list[list[str]] = [
        ["MTD", "mzTab-version", "2.0.0-M"],
        ["MTD", "mzTab-ID", destination.stem],
        ["MTD", "software[1]", "[,,RIKEN Lipidomics Excel to mzTab-M converter,1.0]"],
        ["MTD", "small_molecule-quantification_unit", f"[,,{parsed.abundance_kind} abundance,]"],
        ["MTD", "comment[1]", f"Source workbook: {parsed.source.name}; sheet: {parsed.sheet}"],
    ]
    for assay_index, (name, metadata) in enumerate(zip(parsed.sample_names, parsed.sample_metadata), 1):
        lines.append(["MTD", f"assay[{assay_index}]", clean(name)])
        for custom_index, label in enumerate(sorted(metadata), 1):
            lines.append([
                "MTD",
                f"assay[{assay_index}]-custom[{custom_index}]",
                f"[,,{label},{clean(metadata[label])}]",
            ])
    for variable_index, (label, assay_indices) in enumerate(groups.items(), 1):
        refs = "|".join(f"assay[{index}]" for index in assay_indices)
        lines.extend([
            ["MTD", f"study_variable[{variable_index}]", clean(label)],
            ["MTD", f"study_variable[{variable_index}]-assay_refs", refs],
            ["MTD", f"study_variable[{variable_index}]-description", clean(label)],
        ])

    indices = {
        "id": column_index(parsed, "Alignment ID", "Peak ID"),
        "name": column_index(parsed, "Metabolite name", "Name"),
        "formula": column_index(parsed, "Formula", "Chemical formula"),
        "smiles": column_index(parsed, "SMILES"),
        "inchikey": column_index(parsed, "INCHIKEY", "InChIKey"),
        "adduct": column_index(parsed, "Adduct type", "Adduct"),
        "ontology": column_index(parsed, "Ontology", "Lipid class"),
        "mz": column_index(parsed, "Average Mz", "Average m/z", "m/z"),
        "rt": column_index(parsed, "Average Rt(min)", "Average RT", "RT (min)"),
        "msms": column_index(parsed, "MS/MS assigned"),
        "spectrum": column_index(parsed, "MS/MS spectrum"),
    }
    assay_headers = [f"abundance_assay[{index}]" for index in range(1, len(parsed.sample_names) + 1)]
    header = [
        "SMH", "SML_ID", "SMF_ID_REFS", "database_identifier", "chemical_formula",
        "smiles", "inchi", "chemical_name", "uri", "theoretical_neutral_mass",
        "adduct_ions", "reliability", "best_id_confidence_measure",
        "best_id_confidence_value", *assay_headers, "opt_global_Ontology",
        "opt_global_inchikey", "opt_global_average_mz", "opt_global_retention_time_min",
        "opt_global_ms2_presence", "opt_global_msms_spectrum",
    ]
    lines.append(header)
    for sequence, row in enumerate(parsed.rows, 1):
        row_id = clean(value_at(row, indices["id"]), str(sequence))
        abundances = [clean(value_at(row, column)) for column in parsed.abundance_columns]
        lines.append([
            "SML", row_id, row_id, clean(value_at(row, indices["inchikey"])),
            clean(value_at(row, indices["formula"])), clean(value_at(row, indices["smiles"])),
            NULL, clean(value_at(row, indices["name"])), NULL, NULL,
            clean(value_at(row, indices["adduct"])), NULL, NULL, NULL, *abundances,
            clean(value_at(row, indices["ontology"])), clean(value_at(row, indices["inchikey"])),
            clean(value_at(row, indices["mz"])), clean(value_at(row, indices["rt"])),
            clean(value_at(row, indices["msms"])), clean(value_at(row, indices["spectrum"])),
        ])

    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8", newline="\n") as handle:
        for fields in lines:
            handle.write("\t".join(clean(field) for field in fields) + "\n")


def input_files(path: Path) -> Iterable[Path]:
    if path.is_file():
        yield path
    elif path.is_dir():
        yield from sorted(path.glob("*.xlsx"))
    else:
        raise FileNotFoundError(path)


def output_path(
    source: Path, parsed: ParsedSheet, output: Path, multiple_inputs: bool, multiple_sheets: bool
) -> Path:
    if not multiple_inputs and not multiple_sheets and output.suffix.lower() in {".mztab", ".txt"}:
        return output
    suffix = f"__{re.sub(r'[^A-Za-z0-9_.-]+', '_', parsed.sheet)}" if multiple_sheets else ""
    return output / f"{source.stem}{suffix}.mzTab"


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="An .xlsx file or a directory of .xlsx files")
    parser.add_argument("output", type=Path, help="Output .mzTab file or directory")
    parser.add_argument("--sheet", help="Convert only this worksheet")
    parser.add_argument("--abundance", choices=("normalized", "raw"), default="normalized")
    parser.add_argument("--privacy", choices=("public", "all"), default="public",
                        help="Public samples only (default) or all samples")
    parser.add_argument("--group-by", default="Category,Tissue/Species",
                        help="Comma-separated workbook metadata fields for study variables")
    parser.add_argument("--strict", action="store_true", help="Stop at the first invalid workbook")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    sources = list(input_files(args.input))
    if not sources:
        raise SystemExit(f"No .xlsx files found in {args.input}")
    group_by = [key(item) for item in args.group_by.split(",") if key(item)]
    converted = 0
    failures: list[str] = []
    for source in sources:
        try:
            sheets = parse_workbook(source, args.sheet, args.abundance, args.privacy)
            for parsed in sheets:
                destination = output_path(
                    source, parsed, args.output,
                    multiple_inputs=len(sources) > 1,
                    multiple_sheets=len(sheets) > 1,
                )
                write_mztab(parsed, destination, group_by)
                print(f"converted\t{source}\t{parsed.sheet}\t{destination}\t{len(parsed.rows)} features\t{len(parsed.sample_names)} assays")
                converted += 1
        except Exception as exc:
            message = f"failed\t{source}\t{exc}"
            print(message, file=sys.stderr)
            failures.append(message)
            if args.strict:
                return 1
    print(f"summary\tconverted={converted}\tfailed={len(failures)}")
    return 1 if not converted else 0


if __name__ == "__main__":
    raise SystemExit(main())
